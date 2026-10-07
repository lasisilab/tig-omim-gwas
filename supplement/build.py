"""Generate the Word supplement and compact CSV tables from analysis outputs.
Run: python supplement/build.py (requires python-docx).
"""
from pathlib import Path
import csv
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'supplement'

def read(path):
    with (ROOT / path).open(newline='', encoding='utf-8-sig') as f:
        return list(csv.DictReader(f))

def write(name, rows):
    with (OUT / name).open('w', newline='') as f:
        w=csv.DictWriter(f, fieldnames=list(rows[0])); w.writeheader(); w.writerows(rows)

queries=json.loads((ROOT/'gwas/inputs/data/gwas/search-terms.json').read_text())['queries']
terms=read('gwas/outputs/03-selected-terms.csv')
rules=read('gwas/inputs/data/gwas/phenotype-scope-rules.csv')
genes=read('supplementary-tables/euler-genes.csv')
write('table-s1-searches.csv',queries)
write('table-s1-selected-terms.csv',[{k:r[k] for k in ('group','label','efo_id','scope')} for r in terms])
write('table-s1-phenotype-rules.csv',rules)
cols=['trait_group','gene_symbol','euler_region','omim_cited_ids','gwas_pubmed_ids']
compact=[{k:r[k] for k in cols} for r in genes]
write('table-s2-genes.csv',compact)

doc=Document()
sec=doc.sections[0]
sec.page_width=Inches(11.7);sec.page_height=Inches(8.3)
sec.top_margin=sec.bottom_margin=Inches(.6)
sec.left_margin=sec.right_margin=Inches(.65)
for name in ['Normal','Title','Heading 1','Heading 2','Caption']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    st.font.size=Pt(10 if name=='Normal' else 12)
    st.paragraph_format.space_after=Pt(6)
doc.styles['Title'].font.size=Pt(15)
for el in list(doc.styles.element.iter(qn('w:pBdr'))):
    el.getparent().remove(el)
doc.styles['Normal'].paragraph_format.line_spacing=1
footer=sec.footer.paragraphs[0];footer.alignment=2
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)

def table(headers,rows,widths):
    t=doc.add_table(rows=1, cols=len(headers));t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    t._tbl.tblPr.append(borders)
    for c,h,w in zip(t.rows[0].cells,headers,widths):
        c.text=h;c.width=Inches(w)
        shading=OxmlElement('w:shd');shading.set(qn('w:fill'),'EEEEEE');c._tc.get_or_add_tcPr().append(shading)
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        cells=t.add_row().cells
        for c,v,w in zip(cells,row,widths):c.text=str(v);c.width=Inches(w)
    for i,row in enumerate(t.rows):
        no_split=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(no_split)
        for c in row.cells:
            for p in c.paragraphs:
                p.paragraph_format.space_after=Pt(1.5);p.paragraph_format.space_before=Pt(1.5)
                for run in p.runs:run.font.size=Pt(9);run.bold=(i==0)
    return t

doc.add_paragraph('Classical foundations and genomic advances in human pigmentation and hair morphology', 'Title')
doc.add_paragraph('Yemko Pryor, Lily Heald, Tina Lasisi')
doc.add_paragraph('Supplemental information', 'Heading 1')
doc.add_paragraph('Methods and Tables S1–S2 accompany Figure 3.')
doc.add_paragraph('Supplementary methods','Heading 1')
doc.add_paragraph('We extract OMIM IDs from D’Arcy et al. [S1], Supplementary Table S1, and Needle et al. [S2], Table 2. We assign D’Arcy records and Needle hair color and graying records to pigmentation, and the remaining Needle traits to hair. We deduplicate the cited IDs, retrieve entry names through the OMIM entry API (https://api.omim.org/api/entry), resolve moved entries, and retrieve gene mappings with include=geneMap. We retain the returned approved gene symbols and deduplicate them within each trait group.')
doc.add_paragraph('We search the GWAS Catalog v2 API (https://www.ebi.ac.uk/gwas/rest/api/v2) using the terms in Table S1. We classify returned mapped terms as strict or broad, retrieve associations for both categories without child-trait expansion, and apply the reported-phenotype rules in Table S1. We retrieve study and ancestry records and retain mapped gene names supported by at least two distinct nonempty PubMed IDs within each trait group. We apply no sample-size threshold. Separate publications may include overlapping cohorts.')
doc.add_paragraph('We match exact gene names between OMIM and the retained GWAS sets without alias harmonization. We calculate source-specific and shared sets and plot them with eulerr 8.3.1, using one area scale for both traits. Table S2 lists every plotted gene and its source identifiers. Saved API responses, request URLs, retrieval times and checksums preserve the analyzed snapshot; fresh retrievals may change the results.')
doc.add_paragraph('Code and data','Heading 2')
doc.add_paragraph('The tig-omim-gwas repository contains the retrieval scripts, query settings, analysis code and CSV tables. The root reproduce.py runs the analysis from saved responses and generates the Euler plots.')
doc.add_paragraph('Supplementary references','Heading 2')
doc.add_paragraph('S1. D’Arcy C. et al. (2023) Disease-gene networks of skin pigmentation disorders and reconstruction of protein-protein interaction networks. Bioengineering 10, 13. https://doi.org/10.3390/bioengineering10010013')
doc.add_paragraph('S2. Needle C.D. et al. (2025) A comprehensive review of GWASs of human hair traits. J. Invest. Dermatol. 145, 2964–2972. https://doi.org/10.1016/j.jid.2025.07.004')
doc.add_paragraph('Table S1 GWAS Catalog searches and selection','Heading 1').paragraph_format.page_break_before=True
doc.add_paragraph('Queries and selection rules used for Figure 3. Strict denotes core traits; broad denotes related traits. Both categories enter the analysis.')
doc.add_paragraph('Keyword queries','Heading 2')
table(['Trait group','Search terms'],[(g,', '.join(q['text'] for q in queries if q['group']==g)) for g in ['pigmentation','hair']],[1.2,9.2])
doc.add_paragraph('Scope overrides: ' + '; '.join(f"{r['query_group']}: {r['field']} matching {r['pattern']} is {r['scope']}" for r in rules) + '.')
doc.add_paragraph('Selected mapped terms','Heading 2')
table(['Trait group','Catalog label','Ontology ID','Scope'],[(r['group'],r['label'],r['efo_id'],r['scope']) for r in terms],[1.2,5.4,2.6,1.2])
doc.add_paragraph('Table S2 Genes represented in Figure 3','Heading 1').paragraph_format.page_break_before=True
doc.add_paragraph(f'All {len(genes)} gene–trait records. Shared means membership in both selected sets. OMIM IDs are the literature-cited IDs; PubMed IDs identify supporting GWAS papers. Blank cells indicate absence from the corresponding selected set. Full provenance is retained in supplementary-tables/euler-genes.csv.')
table(['Trait','Gene','Membership','OMIM IDs','GWAS PubMed IDs'],[[r[k] for k in cols] for r in compact],[1,1,1.1,2.3,5])
doc.core_properties.title='Supplemental information for pigmentation and hair morphology'
doc.core_properties.author='Yemko Pryor; Lily Heald; Tina Lasisi'
assert [[c.text for c in row.cells] for row in doc.tables[-1].rows[1:]] == [[r[k] for k in cols] for r in compact]
doc.save(OUT/'supplementary-information.docx')
print(f'Wrote supplementary-information.docx and CSV tables; {len(genes)} gene–trait records.')
